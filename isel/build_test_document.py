from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='04_Teste/I-Crash_Test_Workflow_and_Plan.docx'
def shade(c, fill):
    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(shd)
def borders(t):
    tb=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); tb.append(x)
    t._tbl.tblPr.append(tb)
def put(c,s,b=False,color='000000',size=8.5):
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(str(s)); r.bold=b; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(doc,heads,rows):
    t=doc.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
    for i,h in enumerate(heads): shade(t.rows[0].cells[i],'1F4E78'); put(t.rows[0].cells[i],h,True,'FFFFFF')
    for j,row in enumerate(rows):
        cs=t.add_row().cells
        for i,v in enumerate(row):
            if j%2: shade(cs[i],'F2F6F9')
            put(cs[i],v)
    doc.add_paragraph()
def bullet(doc,s): doc.add_paragraph(s,style='List Bullet')

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
doc.styles['Normal'].font.name='Aptos'; doc.styles['Normal'].font.size=Pt(10)
for s in ('Title','Heading 1','Heading 2'):
    doc.styles[s].font.name='Aptos Display'; doc.styles[s].font.color.rgb=RGBColor(0,0,0)
doc.styles['Heading 1'].font.size=Pt(16); doc.styles['Heading 2'].font.size=Pt(12)
p=doc.add_paragraph(style='Title'); p.add_run('I Crash Test Workflow and Plan')
p=doc.add_paragraph(); p.add_run('Purpose: ').bold=True; p.add_run('Translate the current code workflow into a repeatable test document for regression, security, and release readiness.')
p=doc.add_paragraph(); p.add_run('Scope baseline: ').bold=True; p.add_run('Flutter V2 client, Firebase Auth and Firestore emulator/cloud boundary, scanner services, inventory rules, reporting, and the legacy entry point that remains reachable.')
p=doc.add_paragraph(); p.add_run('Assessment: ').bold=True; p.add_run('Core workflows have strong unit, widget, Rules, and live Android QA evidence. The main remaining risk is incomplete operational workflows and limited repeatable coverage for offline, physical-camera, accessibility, performance, and release-signing scenarios.')

doc.add_heading('1 Workflow under test',1)
doc.add_paragraph('The application is a role-aware, institution-scoped workflow. Authentication determines entry; membership determines visible institutions; role and cart assignment determine data and actions. Business-critical writes must be enforced both in the UI and by Firestore Rules, with Rules as the security boundary.')
table(doc,['Stage','Code path / evidence','Primary test question'],[
('1 Login','AuthGate → LoginScreen','Can the user authenticate and reach the correct next state?'),('2 Institution selection','InstitutionSelectionScreen → memberIndex/memberships','Are only permitted institutions shown?'),('3 Dashboard','InstitutionHomeScreen → carts, summaries, activity','Does the role see the correct carts and actions?'),('4 Cart operations','CartDetailScreen and dialogs','Are structural actions manager+ only?'),('5 Drawer and slots','SlotEditorScreen → replaceSlots','Do layout changes persist and clean orphans?'),('6 Inventory','AssignmentDialog → assign/replenish/consume/reconcile/correct','Are quantities, batches, expiry, and immutable history correct?'),('7 Scanning','GS1 Data Matrix and internal QR services','Are payloads safe, cancellable, and permission-respecting?'),('8 History','HistoryScreen, summaries, CSV/PDF builders','Do views and exports match source events?')])

doc.add_heading('2 Test architecture and execution layers',1)
table(doc,['Layer','What it verifies','Evidence / command'],[
('Domain and service unit','Inventory rules, GS1 parser, HID stream, QR payloads, summaries/exports','flutter test'),('Widget and presentation','Navigation, dialogs, role-gated controls, fake repositories','flutter test'),('Firestore Rules','Authentication, institution isolation, list/get semantics, scoped writes, immutability','firebase emulators:exec --only firestore --project demo-icrash-v2 "npm --prefix firestore-tests test"'),('Live functional QA','Android emulator flow, persistence, role switching, UI/runtime errors','ICrash_API_36 with seeded data'),('Build/static','Analyzer, APK/platform builds, environment correctness','flutter analyze; flutter build apk --debug')])
doc.add_paragraph('Evidence rule: a scenario is verified only when its observable UI result and, where relevant, persisted Firestore state agree. A green build alone is not release evidence.')

doc.add_heading('3 Regression test catalogue',1)
table(doc,['ID','Scenario','Setup','Expected result','Priority'],[
('AUTH-01','Valid login','Valid emulator account','Correct post-login screen; no fatal error','P1'),('SEC-01','Institution isolation','User belongs to institution B','Institution A and its carts are absent and unreadable','P0'),('SEC-02','Cart assignment isolation','User assigned to one cart','Only assigned cart appears; other cart is denied','P0'),('SEC-03','Role action matrix','User, manager, admin, super admin','UI actions and Rules permissions match contract','P0'),('CART-01','Duplicate cart','Cart with drawers, slots, stock, history','Structure copies; stock and history do not','P1'),('LAY-01','Merge/split slots','Edit and save drawer layout','Layout persists; layoutVersion increments; no false orphan','P1'),('INV-01','Consume stock','Quantity above zero','Aggregate decreases; batches untouched','P0'),('INV-02','Invalid consumption','Zero, negative, over-available amounts','Clear failure; quantity never negative','P0'),('INV-03','Replenish expiry','Earlier and later new expiry','Earlier may move known expiry earlier; later never moves it later','P0'),('INV-04','Physical reconciliation','Manager confirms batches and quantity','State derives only from physical confirmation','P0'),('INV-05','Correct history','Correct latest event','Original immutable; compensating event appended','P0'),('SCAN-01','GS1 parse/fallback','Valid, malformed, separator, boundary-year payloads','Fields prefill; malformed remains retryable; manual entry works','P1'),('SCAN-02','Internal QR','Valid/invalid scheme and version','Valid cart opens; invalid payload rejected','P1'),('HIST-01','History summaries','Events across products, carts, periods','Filters and summaries agree with source','P1'),('OPS-01','Restart persistence','Complete flow then relaunch','State restores as designed; no duplicate writes','P1'),('OPS-02','Offline/reconnect','Disconnect during read/write','User-visible behavior is clear; replay respects Rules','P1')])

doc.add_heading('4 Critical acceptance workflows',1)
doc.add_heading('4.1 Normal user daily consumption',2)
for s in ['Seed a normal user assigned to exactly one cart and a product with quantity 2.','Log in, select the institution, open the cart, and long-press the product slot.','Confirm consumption and correction are available while replenish, reconcile, edit, duplicate, member, drawer, and responsible-user actions are absent.','Register consumption of 1; inspect visible quantity and persisted assignment.','Read history and confirm the original event is immutable; attempt an unauthorized field update through Rules tests.']:
    bullet(doc,s)
doc.add_paragraph('Pass condition: the user completes the operational action without structural or stock-management privileges, and the data layer records only the allowed mutation.')
doc.add_heading('4.2 Manager replenishment and physical audit',2)
for s in ['Use manual replenishment and the GS1 scan path.','Test valid scan, unknown GTIN, invalid payload, cancellation, and manual correction.','Replenish with earlier and later expiry dates than the current known expiry.','Run reconciliation with a changed batch set and physically confirmed quantity.','Verify batches, quantity, earliestKnownExpiry, usage events, and Rules authorization.']:
    bullet(doc,s)
doc.add_paragraph('Pass condition: scanning is optional; unknown data is not silently attached; conservative expiry behavior is preserved; reconciliation is the only path that replaces physically confirmed batch state.')
doc.add_heading('4.3 Institution administrator membership management',2)
for s in ['Create a member with temporary credentials without signing the administrator out.','Change role and active/disabled status; verify memberIndex synchronization.','Sign in as the new account and confirm institution/cart visibility.','Attempt the same writes as manager and normal user; Rules must deny them.']:
    bullet(doc,s)

doc.add_heading('5 Coverage assessment and gaps',1)
table(doc,['Area','Current evidence','Risk / next test'],[
('Core business rules','Inventory unit tests, 31 Rules tests, live Android QA','Keep as P0 regression suite; rerun after repository/Rules changes.'),('Presentation','Widget tests across login, institutions, carts, slots, scanning, history, products','Add explicit role-matrix assertions for every manager-only action.'),('Physical camera','Scanner abstraction and fake scanner coverage','Real camera/Data Matrix validation remains a device requirement.'),('Offline/reconnection','Connectivity banner exists; broad scenarios not independently exercised','Run disconnect/reconnect matrix and inspect queued writes.'),('Periodic/monthly audit','Guided workflow not implemented','Do not mark as tested until screens, model, and acceptance rules exist.'),('Daily checklist','Not implemented','Track as product gap, not a passing test.'),('Audit log UI','Data layer exists; presentation wiring absent','Test only exposed immutable usage history.'),('Templates','Duplication exists; template gallery absent','Verify duplicate semantics only.'),('Release readiness','Analyze/build evidence exists; signing/environment separation open','Add release configuration and install/upgrade verification.')])

doc.add_heading('6 Execution checklist',1)
for s in ['Confirm branch, commit, Flutter/Dart version, emulator profile, and backend mode.','Start local Auth/Firestore emulators and seed deterministic QA data.','Run unit and widget tests; record command, date, commit, and result.','Run Firestore Rules tests through firebase emulators:exec; preserve full pass count.','Build/install Android debug APK; complete the three critical workflows.','Inspect UI, emulator logs, persisted Firestore documents, and role changes.','Record defects with scenario ID, expected/actual result, severity, evidence, and retest status.','Do not describe explicitly unimplemented features as passing.','Update docs/TEST_STATUS.md with date and exact environment.']:
    bullet(doc,s)
doc.add_heading('7 Traceability references',1)
doc.add_paragraph('Primary sources: docs/ARCHITECTURE.md; docs/TEST_STATUS.md; docs/RELEASE_CHECKLIST.md; docs/ICRASH_V2_SPECIFICATION.md; firestore.rules; firestore-tests/rules.test.mjs; 03_Implementacao/app1/lib/src/domain/inventory_rules.dart; scanner services under 03_Implementacao/app1/lib/src/services; presentation tests under 03_Implementacao/app1/test.')
doc.add_paragraph('Document status: designed from the repository state inspected on 11 October 2026. Update this baseline whenever the workflow, Rules, or implementation status changes.')
doc.save(OUT)
print(OUT)
